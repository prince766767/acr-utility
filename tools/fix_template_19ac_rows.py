"""One-time fix of the entry rows of 19(a) (table 1) and 19(c) (table 2) in ACR_EMPLOYEE_MASTER.docx
(official form, PDF page 2: short rows with plain cells).

- 19(c): rows had a 2.2 cm minimum height from the old teacher's entries; now the same one-line minimum as 19(a).
  Four columns centred their text vertically and the Class column did not; all cells now align to the top,
  like 19(a) and 19(d).
- 19(a) and 19(c): cells kept extra empty paragraphs from the old entries; each cell keeps only its first one.
Refuses to run if the tables are not in the expected state.
"""
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

MASTER = Path(__file__).resolve().parents[1] / 'ACR_EMPLOYEE_MASTER.docx'


def fail(msg):
    raise SystemExit(msg + ' Nothing changed.')


def text(el):
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))


def drop_extra_paragraphs(rows):
    removed = 0
    for tr in rows:
        for tc in tr.tc_lst:
            for p in tc.findall(qn('w:p'))[1:]:
                if text(p).strip():
                    fail('An entry cell has text in a second paragraph.')
                tc.remove(p)
                removed += 1
    return removed


def main():
    doc = Document(MASTER)
    t1, t2 = doc.tables[1], doc.tables[2]
    rows1 = t1._tbl.tr_lst[1:-1]  # the 5 entry rows; the last row is "Total periods per week"
    rows2 = t2._tbl.tr_lst[1:]
    height1 = rows1[0].find(qn('w:trPr') + '/' + qn('w:trHeight')).get(qn('w:val'))
    heights2 = [tr.find(qn('w:trPr') + '/' + qn('w:trHeight')).get(qn('w:val')) for tr in rows2]
    if len(rows1) != 5 or len(rows2) != 4 or set(heights2) == {height1}:
        fail('Tables 1/2 are not in the expected state.')
    if any(text(tc).strip() for tr in rows1 + rows2 for tc in tr.tc_lst):
        fail('Entry rows of 19(a)/19(c) are not empty.')

    for tr in rows2:
        tr.find(qn('w:trPr') + '/' + qn('w:trHeight')).set(qn('w:val'), height1)
        for tc in tr.tc_lst:
            va = tc.tcPr.find(qn('w:vAlign')) if tc.tcPr is not None else None
            if va is not None:
                tc.tcPr.remove(va)
    removed = drop_extra_paragraphs(rows1) + drop_extra_paragraphs(rows2)
    doc.save(MASTER)
    print(f'Fixed 19(a)/19(c) entry rows: 19(c) row height {heights2[0]} -> {height1} twips, top alignment, '
          f'{removed} leftover empty paragraphs removed.')


if __name__ == '__main__':
    main()
