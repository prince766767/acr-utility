"""One-time: append the official instruction pages 28-30 (appendix_pages/page-*.png) to ACR_EMPLOYEE_MASTER.docx as
full-page image sections, with exactly the code generate_acr.py used to append them to every generated file.
The generators then no longer add them (so the app's generator needs no image handling).
Refuses to run if the template already ends with a picture.
"""
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parents[1]
MASTER = HERE / 'ACR_EMPLOYEE_MASTER.docx'
BLIP = '{http://schemas.openxmlformats.org/drawingml/2006/main}blip'


def add_page_break_image(doc, path):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.page_width = Inches(8.27); sec.page_height = Inches(11.69)
    sec.top_margin = Inches(0); sec.bottom_margin = Inches(0); sec.left_margin = Inches(0); sec.right_margin = Inches(0)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(path), width=Inches(8.27), height=Inches(11.69))


def main():
    doc = Document(MASTER)
    paras = [el for el in doc.element.body if el.tag == qn('w:p')]
    if paras and list(paras[-1].iter(BLIP)):
        raise SystemExit('The template already ends with a picture; nothing changed.')
    for n in (28, 29, 30):
        add_page_break_image(doc, HERE / 'appendix_pages' / f'page-{n}.png')
    doc.save(MASTER)
    print('Appended appendix pages 28-30 to the template.')


if __name__ == '__main__':
    main()
