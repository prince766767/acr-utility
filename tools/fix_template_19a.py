"""One-time fix of the point 19(a) table (table 1) in ACR_EMPLOYEE_MASTER.docx:
- remove Word's automatic "1." ... "5." numbering from the Sr. No. cells (the generator writes the numbers;
  the official form, PDF page 2, has a blank column);
- give the restored "Total periods per week" label the cell's own bold formatting (it was added as a plain run)
  and the header's text colour.
Refuses to run if the table is not in the expected state.
"""
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

MASTER = Path(__file__).resolve().parents[1] / 'ACR_EMPLOYEE_MASTER.docx'


def main():
    doc = Document(MASTER)
    t = doc.tables[1]
    numbered = [n for tr in t._tbl.tr_lst[1:] for n in tr.iter(qn('w:numPr'))]
    runs = t.rows[-1].cells[0].paragraphs[0].runs
    if len(numbered) != 5 or len(runs) != 2 or runs[0].text or runs[1].text != 'Total periods per week' or not runs[0].bold:
        raise SystemExit('Table 1 (19(a)) is not in the expected state; nothing changed.')
    for n in numbered:
        n.getparent().remove(n)
    runs[0].text = 'Total periods per week'
    runs[0].font.color.rgb = t.rows[0].cells[0].paragraphs[0].runs[0].font.color.rgb  # form text colour, not entry blue
    runs[1]._r.getparent().remove(runs[1]._r)
    doc.save(MASTER)
    print('Fixed table 1 (19(a)): removed automatic numbering; "Total periods per week" now bold like the form.')


if __name__ == '__main__':
    main()
