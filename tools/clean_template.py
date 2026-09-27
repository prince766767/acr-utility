"""One-time clean-up of ACR_EMPLOYEE_MASTER.docx so that it matches the blank official form.

The master was made from one teacher's filled ACR and still carried that teacher's entries. This script:
  1. blanks every leftover entry: text runs coloured as filled-in data (blue 0000CC, green 006600, red FF0000),
     except {{TOKENS}} the generator replaces;
  2. deletes the leftover "** Besides these courses ..." row from the 26(ii) table (table 9);
  3. removes Word auto-numbering (and the extra empty numbered lines) from the entry rows of tables 7 and 9-12,
     where it printed "1. 1" next to the serial number the generator writes;
  4. restores the missing "E (ii)" serial label in point 44 (table 29);
  5. puts the {{SESSION}} token in the PBAS header (it held "2024-25");
  6. restores two certificate lines typed in the form's own colour (inside text boxes) to the form's wording.
Run with --dry-run to list what would change. Refuses to run on an already-clean template.
"""
import sys
from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

MASTER = Path(__file__).resolve().parents[1] / 'ACR_EMPLOYEE_MASTER.docx'
DATA_COLOURS = {'0000CC', '006600', 'FF0000'}
FORM_COLOUR = '221F20'
# Coloured runs that are really the official form's own wording (UGC_ACR_Form.pdf pages 1 and 12): kept, made black.
KEEP_AS_FORM = {'Date:', 'Place:', 'Principal', 'Govt. Degree', '__________________'}
# Coloured runs where the form's wording was extended with a college name: restored to the form's wording.
RESTORE_FORM = {'College Alpha (Beta)': 'College.'}
# Text nodes in the form's own colour that still carry the teacher's entries (UGC_ACR_Form.pdf page 12 wording).
RESTORE_TEXT = {'Signature (with stamp) of Principal Govt. Degree College Alpha (Beta)': 'Signature (with stamp) of Principal Govt. Degree College.',
                'Designation, Assistant Professor (Botany)': 'Designation,'}


def run_colour(r):
    rpr = r.find(qn('w:rPr'))
    c = rpr.find(qn('w:color')) if rpr is not None else None
    return (c.get(qn('w:val')) or '').upper() if c is not None else ''


def run_text(r):
    return ''.join(t.text or '' for t in r.iter(qn('w:t')))


def set_text(r, text):
    ts = list(r.iter(qn('w:t')))
    for i, t in enumerate(ts):
        t.text = text if i == 0 else ''


def make_form_colour(r):
    r.find(qn('w:rPr')).find(qn('w:color')).set(qn('w:val'), FORM_COLOUR)


def leftover_runs(body):
    return [r for r in body.iter(qn('w:r'))
            if run_colour(r) in DATA_COLOURS and run_text(r).strip() and '{{' not in run_text(r)]


def main(dry):
    doc = Document(MASTER)
    body = doc.element.body
    runs = leftover_runs(body)
    if not runs:
        raise SystemExit('No leftover entries found; the template is already clean. Nothing changed.')
    t9 = doc.tables[9]
    besides = [row for row in t9.rows if row.cells[0].text.strip().startswith('** Besides')]
    if len(besides) != 1:
        raise SystemExit('Table 9: expected exactly one "** Besides" row; nothing changed.')

    kept = [r for r in runs if run_text(r) in KEEP_AS_FORM]
    restored = [r for r in runs if run_text(r) in RESTORE_FORM]
    blanked = [r for r in runs if r not in kept and r not in restored]
    print(f'1a. Keep {len(kept)} runs of form wording (made black):', ' | '.join(run_text(r) for r in kept))
    print(f'1b. Restore form wording:', ' | '.join(f'{run_text(r)} -> {RESTORE_FORM[run_text(r)]}' for r in restored))
    print(f'1c. Blank {len(blanked)} leftover entries:')
    print('   ' + ' | '.join(run_text(r).strip()[:60] for r in blanked))
    for r in kept:
        make_form_colour(r)
    for r in restored:
        set_text(r, RESTORE_FORM[run_text(r)])
        make_form_colour(r)
    for r in blanked:
        set_text(r, '')
    # The PBAS header on page 6 carried "2024-25"; it gets the same {{SESSION}} token page 1 uses.
    pbas = [p for p in body.iter(qn('w:p')) if 'FOR THE SESSION/YEAR' in ''.join(t.text or '' for t in p.iter(qn('w:t')))]
    if len(pbas) != 1:
        raise SystemExit('Expected one "FOR THE SESSION/YEAR" paragraph; nothing changed.')
    session_runs = [r for r in pbas[0].iter(qn('w:r')) if r in blanked]
    if not session_runs:
        raise SystemExit('No session value found in the PBAS header; nothing changed.')
    set_text(session_runs[0], '{{SESSION}}')
    print('1d. PBAS header session -> {{SESSION}}')

    print('2. Delete table 9 row "** Besides these courses ..."')
    t9._tbl.remove(besides[0]._tr)

    removed_num = removed_par = 0
    for ti in (7, 9, 10, 11, 12):
        for tr in doc.tables[ti]._tbl.tr_lst[1:]:
            for tc in tr.tc_lst:
                paras = tc.findall(qn('w:p'))
                for p in paras:
                    ppr = p.find(qn('w:pPr'))
                    num = ppr.find(qn('w:numPr')) if ppr is not None else None
                    if num is not None:
                        ppr.remove(num)
                        removed_num += 1
                for p in paras[1:]:
                    if not ''.join(t.text or '' for t in p.iter(qn('w:t'))).strip():
                        tc.remove(p)
                        removed_par += 1
    print(f'3. Remove auto-numbering from {removed_num} paragraphs and {removed_par} extra empty lines in tables 7, 9-12')

    t29 = doc.tables[29]
    label, ref = t29.rows[8].cells[0], t29.rows[6].cells[0]
    if label.text.strip() != '' or ref.text.strip() != 'E (i)' or not t29.rows[8].cells[4].text.strip().startswith('10 /'):
        raise SystemExit('Table 29: E(ii) label row not as expected; nothing changed.')
    new_p = deepcopy(ref.paragraphs[0]._p)
    rs = new_p.findall(qn('w:r'))
    for r in rs[1:]:
        new_p.remove(r)
    for t in rs[0].iter(qn('w:t')):
        t.text = 'E (ii)'
    old_p = label.paragraphs[0]._p
    old_p.addprevious(new_p)
    old_p.getparent().remove(old_p)
    print('4. Set point 44 serial label: E (ii)')

    restored_text = 0
    for t in body.iter(qn('w:t')):
        if t.text in RESTORE_TEXT:
            t.text = RESTORE_TEXT[t.text]
            restored_text += 1
    if restored_text != 3:
        raise SystemExit(f'Expected 3 certificate text nodes to restore, found {restored_text}; nothing changed.')
    print('6. Restore certificate wording in', restored_text, 'text nodes')

    if dry:
        print('\nDry run: nothing saved.')
    else:
        doc.save(MASTER)
        print('\nSaved', MASTER.name)


if __name__ == '__main__':
    main('--dry-run' in sys.argv)
