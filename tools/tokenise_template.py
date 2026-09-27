"""One-time: put named {{TOKENS}} into the answer spots of the cover page, Part I, Part II and the teacher's
certificate in ACR_EMPLOYEE_MASTER.docx, and repair wording left over from a filled ACR.

Spec: docs/superpowers/specs/2026-09-27-profile-part1-part2-to-word-design.md, section 6.
Every step checks what it finds first. Nothing is saved unless every step succeeds, and the script refuses to run
on a template that already has {{FULL_NAME}}.
"""
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docx.text.run import Run

MASTER = Path(__file__).resolve().parents[1] / 'ACR_EMPLOYEE_MASTER.docx'
TOKEN_RE = re.compile(r'\{\{[A-Z0-9_]+\}\}')
MC = '{http://schemas.openxmlformats.org/markup-compatibility/2006}'

# (how to find the label line, where the answer spot is, token)
#   'same': first blank blue run on the label's own line
#   'next': first blank blue run on one of the next 3 lines
SPOTS = [
    (('contains', 'Appraisal of work and conduct'), 'same', '{{FULL_NAME}}'),
    (('contains', 'Full Name (in Capital letter)'), 'same', '{{FULL_NAME}}'),
    (('startswith', 'Father/'), 'same', '{{FATHER_HUSBAND}}'),
    (('exact', 'Employee Code'), 'same', '{{EMPLOYEE_CODE}}'),
    (('contains', 'Subject for which Appointed'), 'same', '{{SUBJECT}}'),
    (('contains', 'Date of appointment(in College Cadre)'), 'same', '{{APPOINTMENT_DATE}}'),
    (('contains', 'Current Designation'), 'same', '{{DESIGNATION}}'),
    (('contains', 'Present Pay Band with Grade Pay'), 'same', '{{PAY_INFO}}'),
    (('contains', 'Date of Promotion'), 'same', '{{PROMOTION}}'),
    (('startswith', 'Academic'), 'same', '{{ACADEMIC_QUAL}}'),
    (('exact', 'Professional'), 'same', '{{PROFESSIONAL_QUAL}}'),
    (('contains', 'Research Degree'), 'same', '{{RESEARCH_DEGREE}}'),
    (('contains', 'In words'), 'same', '{{DOB_WORDS}}'),
    (('contains', 'Permanent/Quasi-permanent'), 'same', '{{SERVICE_STATUS}}'),
    (('contains', 'College/Colleges in which served'), 'same', '{{COLLEGES_SERVED}}'),
    (('contains', 'Roll no (with session)'), 'same', '{{DEPT_EXAM}}'),
    (('contains', 'Any other major assignment'), 'same', '{{OTHER_ASSIGNMENT}}'),
    (('contains', 'Permanent Address (With Pin code)'), 'same', '{{ADDR1}}'),
    (('contains', 'Mobile No.'), 'same', '{{MOBILE}}'),
    (('contains', 'Email'), 'same', '{{EMAIL}}'),
    (('contains', 'What do you think has been your most important contribution'), 'next', '{{P17}}'),
    (('contains', 'Have you made any contribution in the area of work not assigned'), 'next', '{{P18}}'),
    (('contains', 'Any special effort made to improve class room'), 'next', '{{P19B}}'),
    (('contains', 'Which new books relating to your subject'), 'next', '{{P19F}}'),
    (('contains', 'What are the vital problems of teaching'), 'next', '{{P19G}}'),
    (('contains', 'Are you doing any Research work'), 'next', '{{RESEARCH_YES_NO}}'),
    (('contains', 'Are you satisfied with your present position'), 'same', '{{P24_SATISFIED}}'),
    (('contains', 'If not, do you want to change the profession'), 'next', '{{P24_REASONS}}'),
]
TITLE_PIECES = {'Dr.': ['Dr.'], '/Shri/': ['/', 'Shri', '/'], 'Smt': ['Smt'], '/Kumari': ['/', 'Kumari']}
RELATION_PIECES = {'Father/': ['Father', '/'], 'Husband': ['Husband']}


def fail(msg):
    raise SystemExit(msg + ' Nothing changed.')


def text(el):
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))


def colour(r):
    rpr = r.find(qn('w:rPr'))
    c = rpr.find(qn('w:color')) if rpr is not None else None
    return (c.get(qn('w:val')) or '').upper() if c is not None else ''


def blank_blue_runs(p):
    return [r for r in p.iter(qn('w:r')) if colour(r) == '0000CC' and r.find(qn('w:t')) is not None and not text(r)]


def set_run_text(r, s):
    ts = list(r.iter(qn('w:t')))
    ts[0].text = s
    ts[0].set(qn('xml:space'), 'preserve')
    for t in ts[1:]:
        t.text = ''


def t_el(s):
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = s
    return t


def new_run(rpr, children, blue=False, unstrike=False):
    r = OxmlElement('w:r')
    if rpr is not None:
        r.append(deepcopy(rpr))
    for c in children:
        r.append(c)
    if blue:
        Run(r, None).font.color.rgb = RGBColor(0, 0, 0xCC)
    if unstrike and r.find(qn('w:rPr')) is not None:
        s = r.find(qn('w:rPr')).find(qn('w:strike'))
        if s is not None:
            s.getparent().remove(s)
    return r


def matches(p, how, label):
    t = text(p).strip()
    return {'contains': label in t, 'startswith': t.startswith(label), 'exact': t == label}[how]


def one_para(paras, how, label):
    hits = [i for i, p in enumerate(paras) if matches(p, how, label)]
    if len(hits) != 1:
        fail(f'Expected exactly one line matching {how} "{label}", found {len(hits)}.')
    return hits[0]


def put_token_in_spots(paras):
    for (how, label), where, token in SPOTS:
        i = one_para(paras, how, label)
        if where == 'next':
            for j in range(i + 1, min(i + 4, len(paras))):
                if blank_blue_runs(paras[j]):
                    i = j
                    break
            else:
                fail(f'No blank answer spot within 3 lines after "{label}".')
        spots = blank_blue_runs(paras[i])
        if not spots:
            fail(f'No blank answer spot on the line of "{label}".')
        set_run_text(spots[0], token)


def mirror_text_box_tokens(body):
    """Word keeps each text box twice: the copy it shows (mc:Choice) and a fallback copy for older readers
    (mc:Fallback). Put the same tokens into the fallback copy's blank answer spots, in the same order."""
    count = 0
    for ac in body.iter(MC + 'AlternateContent'):
        choice, fallback = ac.find(MC + 'Choice'), ac.find(MC + 'Fallback')
        if choice is None or fallback is None:
            continue
        tokens = [m for t in choice.iter(qn('w:t')) for m in TOKEN_RE.findall(t.text or '')]
        if not tokens:
            continue
        spots = [r for r in fallback.iter(qn('w:r')) if colour(r) == '0000CC' and r.find(qn('w:t')) is not None and not text(r)]
        if len(spots) < len(tokens):
            fail(f'Text box with {tokens}: its fallback copy has only {len(spots)} blank answer spots.')
        for tok, r in zip(tokens, spots):
            set_run_text(r, tok)
            count += 1
    return count


def cover_college(paras):
    i = one_para(paras, 'contains', 'Name of the College through which ACR is submitted')
    p = paras[i]
    spots = blank_blue_runs(p)
    after_break = None
    seen_br = False
    for el in p.iter():
        if el.tag == qn('w:br'):
            seen_br = True
        elif el.tag == qn('w:r') and seen_br and el in spots:
            after_break = el
            break
    if not spots or after_break is None or after_break is spots[0]:
        fail('Cover page college line: expected a blank spot before and after the line break.')
    set_run_text(spots[0], '{{COLLEGE_NAME}}')
    set_run_text(after_break, '{{COLLEGE_PLACE}}')


def replace_single_t(body, test, new, what):
    hits = [t for t in body.iter(qn('w:t')) if test(t.text or '')]
    if len(hits) != 1:
        fail(f'Expected exactly one text piece for {what}, found {len(hits)}.')
    hits[0].text = new(hits[0].text)
    hits[0].set(qn('xml:space'), 'preserve')


def restructure(paras, how, label, pieces):
    p = paras[one_para(paras, how, label)]
    runs = [r for r in p.findall(qn('w:r')) if text(r) in pieces]
    if sorted(text(r) for r in runs) != sorted(pieces):
        fail(f'Line "{label}": expected runs {sorted(pieces)}, found {sorted(text(r) for r in runs)}.')
    for r in runs:
        rpr = r.find(qn('w:rPr'))
        for piece in pieces[text(r)]:
            r.addprevious(new_run(rpr, [t_el(piece)], unstrike=True))
        p.remove(r)


def isolate_tokens(body):
    """Split runs so that every {{TOKEN}} is alone in its own blue run (other text keeps its formatting)."""
    while True:
        for r in body.iter(qn('w:r')):
            kids = [c for c in r if c.tag != qn('w:rPr')]
            hit = next(((i, k) for i, k in enumerate(kids) if k.tag == qn('w:t') and TOKEN_RE.search(k.text or '')), None)
            if hit is None:
                continue
            i, k = hit
            m = TOKEN_RE.search(k.text)
            if len(kids) == 1 and k.text == m.group(0) and colour(r) == '0000CC':
                continue
            rpr = r.find(qn('w:rPr'))
            before = kids[:i] + ([t_el(k.text[:m.start()])] if m.start() > 0 else [])
            after = ([t_el(k.text[m.end():])] if m.end() < len(k.text) else []) + kids[i + 1:]
            if before:
                r.addprevious(new_run(rpr, before))
            r.addprevious(new_run(rpr, [t_el(m.group(0))], blue=True))
            if after:
                r.addprevious(new_run(rpr, after))
            r.getparent().remove(r)
            break
        else:
            return


def main():
    doc = Document(MASTER)
    body = doc.element.body
    if '{{FULL_NAME}}' in text(body):
        fail('The template already has {{FULL_NAME}}; it has been tokenised before.')
    paras = list(body.iter(qn('w:p')))

    # Point 25's spot wrongly holds {{OTHER_ASSIGNMENT}} (point 14's content): give it its own token first.
    i25 = one_para(paras, 'contains', 'Any other significant point which is not covered')
    t25 = [t for p in paras[i25 + 1:i25 + 3] for t in p.iter(qn('w:t')) if (t.text or '').strip() == '{{OTHER_ASSIGNMENT}}']
    if len(t25) != 1:
        fail('Point 25: expected {{OTHER_ASSIGNMENT}} in its answer spot.')
    t25[0].text = '{{P25}}'

    cover_college(paras)
    put_token_in_spots(paras)
    if mirror_text_box_tokens(body) != 3:
        fail('Expected the tokens of points 17, 18 and 19(b) in the fallback copies of their text boxes.')

    lp = paras[one_para(paras, 'contains', 'Land line telephone')]
    lines = [r for r in lp.findall(qn('w:r')) if text(r) == '__________________']
    if len(lines) != 1:
        fail('Point 16: expected the ______ line after "Land line telephone No.:".')
    set_run_text(lines[0], '{{LANDLINE}}')

    replace_single_t(body, lambda s: s.startswith('Exempted vide Director of Higher Education letter no.:') and '{{HINDI_DETAILS}}' in s,
                     lambda s: 'b) Hindi subject : Cleared / exempted (mention details) {{HINDI_DETAILS}}', 'point 13(b)')
    replace_single_t(body, lambda s: '{{OTHER_INFO}}' in s, lambda s: s.replace('{{OTHER_INFO}}', ''), 'the unused {{OTHER_INFO}}')
    dp = paras[one_para(paras, 'contains', '{{REPORT_DATE}}')]
    des = [t for t in dp.iter(qn('w:t')) if (t.text or '') == 'Designation,']
    if len(des) != 1:
        fail('Teacher\'s certificate: expected "Designation," on the date line.')
    des[0].text = 'Designation, {{CERT_DESIGNATION}}'
    des[0].set(qn('xml:space'), 'preserve')

    restructure(paras, 'contains', 'Appraisal of work and conduct', TITLE_PIECES)
    restructure(paras, 'startswith', 'Father/', RELATION_PIECES)

    last = doc.tables[1].rows[-1]
    if last.cells[0].text.strip() or len(doc.tables[1].rows) != 7:
        fail('19(a): expected an empty label cell in the last row of table 1.')
    last.cells[0].paragraphs[0].add_run('Total periods per week')

    isolate_tokens(body)
    doc.save(MASTER)
    print('Tokenised the template: cover page, points 1-25, point 16 landline, 13(b) wording, certificate designation,')
    print('unstruck title/relation options, 19(a) total label.')


if __name__ == '__main__':
    main()
