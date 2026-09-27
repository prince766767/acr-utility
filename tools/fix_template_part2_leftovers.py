"""One-time fix of Part II leftovers in ACR_EMPLOYEE_MASTER.docx found when checking a generated ACR against the
official form (PDF pages 2-4):

1. 19(c) (table 2) and 20 (table 4): cells were merged down across the empty entry rows; every row gets its own cells.
2. The questions "Which new books ..." and "What are the vital problems ..." were auto-lettered (e) and (f);
   the form has (f) and (g). The letters become typed text in the same position.
3. The answers to 19(f) and 19(g) carried the old teacher's list numbering ("(i)/(ii)", "(1)-(4)"); removed,
   with its hanging first-line indent.
4. The answers to 19(f), 23, 24 (Yes/No) and 25 were bold (old formatting); made regular like the other answers.
5. Six answer spots kept the old teacher's condensed character spacing; answers now print at normal spacing.
6. The point 8 answer spot was in italics (old "No Promotion yet"); made upright like the other answers.
Refuses to run if the template is not in the expected state.
"""
from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

MASTER = Path(__file__).resolve().parents[1] / 'ACR_EMPLOYEE_MASTER.docx'
LETTERS = {'Which new books relating to your subject': '(f)', 'What are the vital problems of teaching': '(g)'}
UNBOLD = ('{{P19F}}', '{{P23}}', '{{P24_SATISFIED}}', '{{P25}}')
UNCONDENSE = ('{{DOB_WORDS}}', '{{HINDI_DETAILS}}', '{{P17}}', '{{P18}}', '{{P19F}}', '{{P19G}}')


def fail(msg):
    raise SystemExit(msg + ' Nothing changed.')


def text(el):
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))


def num_pr(p):
    ppr = p.find(qn('w:pPr'))
    return ppr.find(qn('w:numPr')) if ppr is not None else None


def unmerge(table, first_row):
    count = 0
    for tr in table._tbl.tr_lst[first_row:]:
        for tc in tr.tc_lst:
            vm = tc.tcPr.find(qn('w:vMerge')) if tc.tcPr is not None else None
            if vm is not None:
                if text(tc).strip():
                    fail('A merged entry cell is not empty.')
                tc.tcPr.remove(vm)
                count += 1
    return count


def main():
    doc = Document(MASTER)
    body = doc.element.body
    paras = list(body.iterchildren(qn('w:p')))

    merged = unmerge(doc.tables[2], 1), unmerge(doc.tables[4], 3)
    if merged != (8, 6):
        fail(f'Expected 8 merged cells in 19(c) and 6 in point 20, found {merged}.')

    for start, letter in LETTERS.items():
        hits = [p for p in paras if text(p).startswith(start)]
        if len(hits) != 1 or num_pr(hits[0]) is None:
            fail(f'Expected one auto-lettered question starting "{start}".')
        p = hits[0]
        p.find(qn('w:pPr')).remove(num_pr(p))
        first = p.find(qn('w:r'))
        r = OxmlElement('w:r')
        if first.find(qn('w:rPr')) is not None:
            r.append(deepcopy(first.find(qn('w:rPr'))))
        t = OxmlElement('w:t')
        t.text = letter
        r.append(t)
        r.append(OxmlElement('w:tab'))
        first.addprevious(r)

    # Answer areas of 19(f) and 19(g): from the answer token down to the next question, drop the old list numbering.
    for token, stop in (('{{P19F}}', 'What are the vital problems'), ('{{P19G}}', 'Details of Last year')):
        idx = [i for i, p in enumerate(paras) if token in text(p)]
        if len(idx) != 1:
            fail(f'Expected one paragraph with {token}.')
        removed = 0
        for p in paras[idx[0]:]:
            if text(p).startswith('(g)') or text(p).startswith(stop):
                break
            n = num_pr(p)
            if n is not None:
                ppr = n.getparent()
                ppr.remove(n)
                # The List Paragraph style hangs the first line out for the number; with no number it would stick
                # out, so the first line starts where the other lines do.
                ind = ppr.find(qn('w:ind'))
                if ind is None:
                    ind = OxmlElement('w:ind')
                    after = ppr.find(qn('w:spacing'))
                    (after.addnext(ind) if after is not None else ppr.find(qn('w:pStyle')).addnext(ind))
                    ind.set(qn('w:left'), str(int(doc.styles['List Paragraph'].paragraph_format.left_indent.twips)))
                if ind.get(qn('w:hanging')) is not None:
                    del ind.attrib[qn('w:hanging')]
                ind.set(qn('w:firstLine'), '0')
                removed += 1
        if removed == 0:
            fail(f'No list numbering found in the answer area of {token}.')

    for token in UNBOLD:
        runs = [r for r in body.iter(qn('w:r')) if text(r) == token]
        if len(runs) != 1:
            fail(f'Expected one run holding {token}.')
        rpr = runs[0].find(qn('w:rPr'))
        for tag in ('w:b', 'w:bCs'):
            b = rpr.find(qn(tag)) if rpr is not None else None
            if b is not None:
                rpr.remove(b)

    runs = [r for r in body.iter(qn('w:r')) if text(r) == '{{PROMOTION}}']
    italic = runs[0].find(qn('w:rPr') + '/' + qn('w:i')) if len(runs) == 1 else None
    if italic is None:
        fail('Expected the point 8 answer spot in italics.')
    italic.getparent().remove(italic)
    ics = runs[0].find(qn('w:rPr') + '/' + qn('w:iCs'))
    if ics is not None:
        ics.getparent().remove(ics)

    for token in UNCONDENSE:
        runs = [r for r in body.iter(qn('w:r')) if text(r) == token]
        rpr = runs[0].find(qn('w:rPr')) if len(runs) == 1 else None
        sp = rpr.find(qn('w:spacing')) if rpr is not None else None
        if sp is None:
            fail(f'Expected condensed character spacing on {token}.')
        rpr.remove(sp)

    doc.save(MASTER)
    print('Fixed Part II leftovers: merged cells in 19(c)/20, question letters (f)/(g), old answer numbering and indent,')
    print('bold answers, condensed answer spacing.')


if __name__ == '__main__':
    main()
