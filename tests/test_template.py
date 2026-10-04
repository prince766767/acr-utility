import re, sys, unittest
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
OLD_TEACHER_WORDS = {'13ff79d7bc39bdba', '2e9dfdfa68634c18', '567af38c766a7689', '5d0569ce4c2a92ab', '62db5835117e3241', '8f00d16bb89ea112', '9bd2ef7a9ddbb223', 'acee714c1985a2e9', 'c067794a648ca326', 'e994b9330ad9f032', 'ed96d148b84152b5', 'f97438b042886e91'}


class Template27iii(unittest.TestCase):
    def test_27iii_matches_official_form(self):
        t = Document(ROOT / 'ACR_EMPLOYEE_MASTER.docx').tables[12]
        col1 = [r.cells[1].text.strip() for r in t.rows]
        self.assertEqual(len(t.rows), 14)
        self.assertEqual(col1[9], '(iii) Professional Development Activities')
        self.assertEqual([t.rows[9].cells[c].text.strip() for c in (0, 2, 3)], ['', '', ''])
        for r in (10, 11):
            self.assertEqual([c.text.strip() for c in t.rows[r].cells], ['', '', '', ''])
        self.assertEqual(col1[12], 'Total (Max.15)')
        self.assertTrue(col1[13].startswith('Total Score (I+II+III)'))
        self.assertNotIn('Chaired', ' '.join(col1))
        heading = [(r.font.color.rgb, r.bold) for r in t.rows[9].cells[1].paragraphs[0].runs[:2]]
        reference = [(r.font.color.rgb, r.bold) for r in t.rows[5].cells[1].paragraphs[0].runs[:2]]
        self.assertEqual(heading, reference)

    def test_entry_rows_have_no_merged_cells(self):
        # Every entry row of the point 26-28 tables needs its own cells, or one score overwrites another.
        from docx.oxml.ns import qn
        doc = Document(ROOT / 'ACR_EMPLOYEE_MASTER.docx')
        for ti in (10, 11, 12, 13, 14, 15, 16, 17, 18, 19, 20, 21, 22):
            for r, tr in enumerate(doc.tables[ti]._tbl.tr_lst):
                for tc in tr.tc_lst:
                    self.assertIsNone(tc.tcPr.find(qn('w:vMerge')) if tc.tcPr is not None else None, f'table {ti} row {r}')


class TemplateIsBlank(unittest.TestCase):
    """The master must match the blank official form: no teacher's leftover entries."""
    @classmethod
    def setUpClass(cls):
        cls.doc = Document(ROOT / 'ACR_EMPLOYEE_MASTER.docx')

    def test_no_leftover_entries(self):
        from docx.oxml.ns import qn
        left = []
        for r in self.doc.element.body.iter(qn('w:r')):
            rpr = r.find(qn('w:rPr'))
            c = rpr.find(qn('w:color')) if rpr is not None else None
            text = ''.join(t.text or '' for t in r.iter(qn('w:t'))).strip()
            if c is not None and (c.get(qn('w:val')) or '').upper() in ('0000CC', '006600', 'FF0000') and text and '{{' not in text:
                left.append(text)
        self.assertEqual(left, [])
        full = ' '.join(t.text or '' for t in self.doc.element.body.iter(qn('w:t')))  # paragraphs and tables
        for s in ('NIL', '2024-25'):
            self.assertNotIn(s, full)
        # Names, places and numbers of the teacher whose filled ACR the master was made from, stored only as hashes
        # so that this test does not itself keep the details.
        import hashlib
        words = {w.lower() for w in re.findall(r'[A-Za-z0-9]+', full)}
        self.assertEqual({w for w in words if hashlib.sha256(w.encode()).hexdigest()[:16] in OLD_TEACHER_WORDS}, set())

    def test_file_properties_blank(self):
        props = self.doc.core_properties
        self.assertEqual((props.author, props.last_modified_by, props.title), ('', '', ''))

    def test_no_auto_numbering_in_entry_rows(self):
        from docx.oxml.ns import qn
        for ti in (0, 1, 2, 3, 4, 5, 6, 24) + tuple(range(7, 23)):
            for tr in self.doc.tables[ti]._tbl.tr_lst[1:]:
                self.assertEqual(tr.findall('.//' + qn('w:numPr')), [], f'table {ti}')

    def test_26ii_has_no_leftover_note_and_44_has_E_ii_label(self):
        self.assertFalse(any('Besides' in r.cells[0].text for r in self.doc.tables[9].rows))
        self.assertEqual(self.doc.tables[29].rows[8].cells[0].text.strip(), 'E (ii)')


class TemplatePage11(unittest.TestCase):
    def test_certificate_lines_match_form(self):
        doc = Document(ROOT / 'ACR_EMPLOYEE_MASTER.docx')
        texts = [p.text for p in doc.paragraphs]
        self.assertIn('Place: {{PLACE}}Signature of the reported on officer', [t.replace('	', '') for t in texts])
        self.assertFalse(any('{{COLLEGE_PIN}}, {{COLLEGE_NAME}}' in t for t in texts))
        sig = [t for t in texts if 'Signature (with stamp)' in t]
        self.assertEqual(len(sig), 1)
        self.assertEqual(sig[0].replace('	', '').strip(), 'Date:Signature (with stamp) of Principal')


class TemplateTokens(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from docx.oxml.ns import qn
        cls.qn = staticmethod(qn)
        cls.doc = Document(ROOT / 'ACR_EMPLOYEE_MASTER.docx')

    def runs_text(self, r):
        return ''.join(t.text or '' for t in r.iter(self.qn('w:t')))

    def test_every_token_alone_in_a_blue_run(self):
        import re
        sys.path.insert(0, str(ROOT))
        from acr_fields import TOKENS
        qn = self.qn
        found = {}
        for r in self.doc.element.body.iter(qn('w:r')):
            if r.find('.//' + qn('w:r')) is not None:
                continue  # a run holding a whole text box; its inner runs are checked themselves
            text = self.runs_text(r)
            for tok in re.findall(r'\{\{([A-Z0-9_]+)\}\}', text):
                found[tok] = found.get(tok, 0) + 1
                rpr = r.find(qn('w:rPr'))
                colour = rpr.find(qn('w:color')).get(qn('w:val')) if rpr is not None and rpr.find(qn('w:color')) is not None else ''
                self.assertEqual((text, colour.upper()), ('{{%s}}' % tok, '0000CC'), tok)
        expected = {t: 1 for t in TOKENS}
        expected['FULL_NAME'] = 2
        expected['SESSION'] = 2
        self.assertEqual(found, expected)

    def test_title_and_relation_options_unstruck_and_separate(self):
        qn = self.qn
        for start, options in (('Appraisal of work and conduct', ('Dr.', 'Shri', 'Smt', 'Kumari')), ('Father/Husband', ('Father', 'Husband'))):
            paras = [p for p in self.doc.element.body.iter(qn('w:p')) if ''.join(t.text or '' for t in p.iter(qn('w:t'))).startswith(start)]
            self.assertEqual(len(paras), 1, start)
            runs = {self.runs_text(r): r for r in paras[0].findall(qn('w:r'))}
            for opt in options:
                self.assertIn(opt, runs, opt)
            for r in paras[0].findall(qn('w:r')):
                rpr = r.find(qn('w:rPr'))
                strike = rpr.find(qn('w:strike')) if rpr is not None else None
                self.assertTrue(strike is None or strike.get(qn('w:val')) in ('0', 'false'), self.runs_text(r))

    def test_wording_repairs(self):
        full = ' '.join(t.text or '' for t in self.doc.element.body.iter(self.qn('w:t')))
        self.assertIn('b) Hindi subject : Cleared / exempted (mention details)', full)
        self.assertNotIn('Exempted vide Director', full)
        self.assertNotIn('{{OTHER_INFO}}', full)
        self.assertEqual(self.doc.tables[1].rows[-1].cells[0].text.strip(), 'Total periods per week')
        self.assertTrue(all(r.bold for r in self.doc.tables[1].rows[-1].cells[0].paragraphs[0].runs if r.text))

    def test_answer_boxes_are_bordered_paragraphs(self):
        # 17, 18 and 19(b): bordered body paragraphs (they grow and split across pages), not text boxes or tables
        qn = self.qn
        for tok in ('{{P17}}', '{{P18}}', '{{P19B}}'):
            ts = [t for t in self.doc.element.body.iter(qn('w:t')) if t.text == tok]
            self.assertEqual(len(ts), 1, tok)
            p = next(ts[0].iterancestors(qn('w:p')))
            self.assertIs(p.getparent(), self.doc.element.body, tok)
            self.assertIsNotNone(p.find(qn('w:pPr') + '/' + qn('w:pBdr')), tok)
        compat = self.doc.settings.element.find(qn('w:compat'))
        self.assertIsNotNone(compat.find(qn('w:doNotExpandShiftReturn')))

    def test_part2_leftovers_fixed(self):
        qn = self.qn
        body = self.doc.element.body
        first = {1: 1, 2: 1, 3: 1, 4: 3, 5: 1, 6: 1, 24: 1}
        for ti, start in first.items():
            for tr in self.doc.tables[ti]._tbl.tr_lst[start:]:
                self.assertEqual([tc for tc in tr.tc_lst if tc.tcPr is not None and tc.tcPr.find(qn('w:vMerge')) is not None], [], f'table {ti}')
        texts = [''.join(t.text or '' for t in p.iter(qn('w:t'))) for p in body.iterchildren(qn('w:p'))]
        self.assertTrue(any(t.startswith('(f)Which new books relating to your subject') for t in texts))
        self.assertTrue(any(t.startswith('(g)What are the vital problems of teaching') for t in texts))
        paras = list(body.iterchildren(qn('w:p')))
        i = next(i for i, p in enumerate(paras) if '{{P19F}}' in ''.join(t.text or '' for t in p.iter(qn('w:t'))))
        j = next(i for i, p in enumerate(paras) if ''.join(t.text or '' for t in p.iter(qn('w:t'))).startswith('Details of Last year'))
        for p in paras[i:j]:
            t = ''.join(x.text or '' for x in p.iter(qn('w:t')))
            if t.startswith('(g)'):
                continue
            self.assertIsNone(p.find(qn('w:pPr') + '/' + qn('w:numPr')), t)
            ind = p.find(qn('w:pPr') + '/' + qn('w:ind'))
            if t.strip():
                self.assertIsNotNone(ind, t)
                self.assertIsNone(ind.get(qn('w:hanging')), t)
                self.assertEqual(ind.get(qn('w:firstLine')), '0', t)
        for r in body.iter(qn('w:r')):
            t = ''.join(x.text or '' for x in r.iter(qn('w:t')))
            if t in ('{{P19F}}', '{{P23}}', '{{P24_SATISFIED}}', '{{P25}}'):
                self.assertIsNone(r.find(qn('w:rPr') + '/' + qn('w:b')), t)
            if t.startswith('{{'):
                self.assertIsNone(r.find(qn('w:rPr') + '/' + qn('w:spacing')), t)
            if t == '{{PROMOTION}}':
                self.assertIsNone(r.find(qn('w:rPr') + '/' + qn('w:i')), t)

    def test_19a_19c_entry_rows_plain_and_short(self):
        qn = self.qn
        t1, t2 = self.doc.tables[1], self.doc.tables[2]
        height = t1._tbl.tr_lst[1].find(qn('w:trPr') + '/' + qn('w:trHeight')).get(qn('w:val'))
        for tr in t1._tbl.tr_lst[1:-1] + t2._tbl.tr_lst[1:]:
            self.assertEqual(tr.find(qn('w:trPr') + '/' + qn('w:trHeight')).get(qn('w:val')), height)
            for tc in tr.tc_lst:
                self.assertEqual(len(tc.findall(qn('w:p'))), 1)
                self.assertIsNone(tc.tcPr.find(qn('w:vAlign')) if tc.tcPr is not None else None)

    def test_point_20_rows_line_up_with_header(self):
        qn = self.qn
        def spans(tr):
            return [int(tc.tcPr.find(qn('w:gridSpan')).get(qn('w:val'))) if tc.tcPr is not None and tc.tcPr.find(qn('w:gridSpan')) is not None else 1
                    for tc in tr.tc_lst]
        trs = self.doc.tables[4]._tbl.tr_lst
        for tr in trs[3:]:
            self.assertEqual(spans(tr), spans(trs[2]))


class TemplateAppendixPages(unittest.TestCase):
    def test_ends_with_three_full_page_images(self):
        from docx.oxml.ns import qn
        doc = Document(ROOT / 'ACR_EMPLOYEE_MASTER.docx')
        paras = [el for el in doc.element.body if el.tag == qn('w:p')]
        blip = '{http://schemas.openxmlformats.org/drawingml/2006/main}blip'
        # each page is a section-break paragraph followed by the picture paragraph
        self.assertEqual([len(list(p.iter(blip))) for p in paras[-6:]], [0, 1, 0, 1, 0, 1])
        self.assertAlmostEqual(doc.sections[-1].page_width.inches, 8.27, places=2)
        self.assertEqual(doc.sections[-1].left_margin, 0)


class TemplateGoogleDocsBreaks(unittest.TestCase):
    """Google Docs (used to make the PDF) turns some Word page breaks into blank pages; see tools/fix_template_google_breaks.py."""
    @classmethod
    def setUpClass(cls):
        from docx.oxml.ns import qn
        cls.qn = staticmethod(qn)
        cls.paras = [el for el in Document(ROOT / 'ACR_EMPLOYEE_MASTER.docx').element.body if el.tag == qn('w:p')]

    def page_break_sections(self):
        qn = self.qn
        for i, p in enumerate(self.paras):
            sect = p.find(qn('w:pPr') + '/' + qn('w:sectPr'))
            if sect is None:
                continue
            kind = sect.find(qn('w:type'))
            if kind is None or kind.get(qn('w:val')) != 'continuous':
                yield i, p

    def test_no_new_page_setting_right_after_a_section_break(self):
        qn = self.qn
        for i, _ in self.page_break_sections():
            if i + 1 < len(self.paras):
                self.assertIsNone(self.paras[i + 1].find(qn('w:pPr') + '/' + qn('w:pageBreakBefore')), i)

    def test_empty_section_break_paragraphs_are_minimal(self):
        qn = self.qn
        seen = 0
        for i, p in self.page_break_sections():
            if ''.join(t.text or '' for t in p.iter(qn('w:t'))).strip():
                continue
            seen += 1
            spacing = p.find(qn('w:pPr') + '/' + qn('w:spacing'))
            self.assertIsNotNone(spacing, i)
            self.assertEqual((spacing.get(qn('w:before')), spacing.get(qn('w:after')), spacing.get(qn('w:line')), spacing.get(qn('w:lineRule'))), ('0', '0', '20', 'exact'), i)
            sz = p.find(qn('w:pPr') + '/' + qn('w:rPr') + '/' + qn('w:sz'))
            self.assertEqual(sz.get(qn('w:val')) if sz is not None else None, '2', i)
        self.assertGreater(seen, 10)

    def test_spacer_at_bottom_of_part3_page_is_minimal(self):
        # the page ending with point 34(d) "...Just good enough."; spacer(s) between it and its section break
        qn = self.qn
        texts = [''.join(t.text or '' for t in p.iter(qn('w:t'))) for p in self.paras]
        def ends_page(k):  # only empty paragraphs between it and a section break
            j = k + 1
            while self.paras[j].find(qn('w:pPr') + '/' + qn('w:sectPr')) is None:
                if texts[j].strip() or self.paras[j].getnext() is not self.paras[j + 1]:
                    return False
                j += 1
            return True
        hits = [k for k, t in enumerate(texts) if t.strip().startswith('Just good enough.') and ends_page(k)]
        self.assertEqual(len(hits), 1)
        i = hits[0]
        j = i + 1
        spacers = 0
        while self.paras[j].find(qn('w:pPr') + '/' + qn('w:sectPr')) is None:
            self.assertFalse(texts[j].strip(), j)
            spacing = self.paras[j].find(qn('w:pPr') + '/' + qn('w:spacing'))
            self.assertEqual((spacing.get(qn('w:line')), spacing.get(qn('w:lineRule'))) if spacing is not None else None, ('20', 'exact'), j)
            spacers += 1
            j += 1
        self.assertGreaterEqual(spacers, 1)


class TemplateNoTextBoxesWhereGoogleBreaks(unittest.TestCase):
    """Google Docs redraws Word text boxes in its own way; see tools/fix_template_textboxes.py."""
    @classmethod
    def setUpClass(cls):
        cls.doc = Document(ROOT / 'ACR_EMPLOYEE_MASTER.docx')

    def box_texts(self):
        wps = '{http://schemas.microsoft.com/office/word/2010/wordprocessingShape}txbx'
        from docx.oxml.ns import qn
        return [' '.join(''.join(t.text or '' for t in b.iter(qn('w:t'))).split()) for b in self.doc.element.body.iter(wps)]

    def test_note_45_is_a_bordered_paragraph_before_heading_45(self):
        from docx.oxml.ns import qn
        self.assertFalse([t for t in self.box_texts() if t.startswith('NOTE:- (*), (**)')])
        body = list(self.doc.element.body)
        heading = next(i for i, el in enumerate(body) if el.tag == qn('w:p') and 'SUMMARY OF API SCORES BY PRINCIPAL' in ' '.join(''.join(t.text or '' for t in el.iter(qn('w:t'))).split()))
        p = body[heading - 1]
        self.assertEqual(p.tag, qn('w:p'))
        self.assertEqual(' '.join(''.join(t.text or '' for t in p.iter(qn('w:t'))).split()), 'NOTE:- (*), (**).. Please refer to clarification mentioned on page7950 of UGC Regulations (THE GAZETTE OF INDIA, September 18, 2010).')
        bdr = p.find(qn('w:pPr') + '/' + qn('w:pBdr'))
        self.assertEqual(sorted(c.tag.split('}')[1] for c in bdr), ['bottom', 'left', 'right', 'top'])

    def test_part_iv_boxes_are_bordered_paragraphs(self):
        from docx.oxml.ns import qn
        for start in ('Teacher Name', 'Overall performance in percentage', 'Note:-'):
            self.assertFalse([t for t in self.box_texts() if t.startswith(start)], start)
            paras = [p for p in self.doc.element.body if p.tag == qn('w:p')
                     and ' '.join(''.join(t.text or '' for t in p.iter(qn('w:t'))).split()).startswith(start)]
            self.assertEqual(len(paras), 1, start)
            self.assertIsNotNone(paras[0].find(qn('w:pPr') + '/' + qn('w:pBdr')), start)

    def test_table_count_unchanged(self):
        # the generators address tables by position
        self.assertEqual(len(self.doc.tables), 44)

if __name__ == '__main__':
    unittest.main()
