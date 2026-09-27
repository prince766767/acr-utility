import json, shutil, sys, tempfile, unittest
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import generate_acr  # noqa: E402

CASES = json.loads((ROOT / 'tests' / 'fixtures' / 'cases.json').read_text(encoding='utf-8'))
WORKED = CASES[0]['api']


def txt(doc, t, r, c):
    return doc.tables[t].rows[r].cells[c].text.strip()


class GenerateApi(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp())
        out = cls.tmp / 'worked.docx'
        generate_acr.generate({'session': '2025-26', 'api': WORKED}, out)
        cls.doc = Document(out)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_session_in_pbas_header(self):
        paras = [p.text for p in self.doc.paragraphs if 'FOR THE SESSION/YEAR' in p.text]
        self.assertEqual(len(paras), 1)
        self.assertIn('2025-26', paras[0])

    def test_point_29_col4_and_col3_untouched(self):
        d = self.doc
        self.assertEqual([txt(d, 23, r, 3) for r in (1, 2, 3, 4)], ['106.75', '25', '131.75', '193'])
        self.assertEqual([txt(d, 23, r, 2) for r in (1, 2, 3, 4)], ['', '', '', ''])

    def test_point_42_col4(self):
        self.assertEqual([txt(self.doc, 25, r, 3) for r in range(3, 9)], ['45', '6', '18.5', '15.25', '22', '106.75'])

    def test_point_43_col4(self):
        self.assertEqual([txt(self.doc, 26, r, 3) for r in range(3, 7)], ['15', '12.5', '3', '25'])

    def test_point_44_col5(self):
        d = self.doc
        expected = CASES[0]['values']['p44']
        for code, (t, r, c, _) in generate_acr.P44_CELLS.items():
            self.assertEqual(txt(d, t, r, c), expected[code], code)
        self.assertEqual(txt(d, 30, 2, 4), '193')

    def test_principal_columns_empty(self):
        d = self.doc
        for t, first, api_col in ((25, 3, 3), (26, 3, 3), (27, 3, 4), (28, 0, 4), (29, 0, 5), (30, 0, 4)):
            for r in range(first, len(d.tables[t].rows)):
                seen = set()
                for cell in d.tables[t].rows[r].cells[api_col + 1:]:
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    self.assertEqual(cell.text.strip(), '', f'table {t} row {r}')

    def test_point_26_tables(self):
        d = self.doc
        self.assertEqual(txt(d, 8, 1, 2), '45')
        self.assertEqual(txt(d, 8, 2, 2), '6')
        self.assertEqual([txt(d, 9, 1, c) for c in range(5)], ['1', 'B.A. I Economics', 'Text book', 'Yes', 'Notes'])
        self.assertEqual(txt(d, 9, 5, 4), '18.5')
        self.assertEqual([txt(d, 10, r, 2) for r in (1, 2, 3)], ['8', '7.25', '15.25'])
        self.assertEqual([txt(d, 11, r, 4) for r in (1, 2, 3)], ['10', '12', '22'])

    def test_point_27_table(self):
        d = self.doc
        self.assertEqual([txt(d, 12, 2, c) for c in range(4)], ['1', 'NSS', '2', '10'])
        self.assertEqual(txt(d, 12, 4, 3), '15')
        self.assertEqual([txt(d, 12, r, 3) for r in (6, 7, 8)], ['7.5', '5', '12.5'])
        self.assertEqual([txt(d, 12, 10, c) for c in range(4)], ['1', 'Seminar organised', 'one-day', '3'])
        self.assertEqual(txt(d, 12, 12, 3), '3')
        self.assertEqual(txt(d, 12, 13, 3), '25')

    def test_point_28_tables(self):
        d = self.doc
        self.assertEqual(len(d.tables[13].rows), 4)  # header + 3 papers
        self.assertEqual([txt(d, 13, r, 7) for r in (1, 2, 3)], ['15', '7.5', '10'])
        self.assertEqual(txt(d, 15, 1, 6), '10')
        self.assertEqual(txt(d, 19, 1, 4), '3')
        self.assertEqual(txt(d, 19, 2, 4), '17')
        self.assertEqual([txt(d, 21, 1, 4), txt(d, 21, 3, 4), txt(d, 21, 4, 4)], ['National', 'Regional / State', 'Local - University / College'])
        self.assertEqual(txt(d, 22, 1, 5), '5')

    def test_more_entries_than_template_rows(self):
        api = {'c1': {'innovative': [{'description': f'm{i}', 'score': 3} for i in range(4)],
                      'exam': [{'type': f'd{i}', 'score': 5} for i in range(5)],
                      'resources': [{'course': f'c{i}'} for i in range(5)], 'resourcesScore': 12},
               'c2': {'extension': [{'activity': f'e{i}', 'score': 4} for i in range(4)],
                      'professional': [{'activity': f'p{i}', 'score': 2.5} for i in range(3)]},
               'c3': {'papers': [{'title': f't{i}', 'row': 'E2d', 'score': 3} for i in range(6)]}}
        out = self.tmp / 'many.docx'
        generate_acr.generate({'api': api}, out)
        d = Document(out)
        self.assertEqual([txt(d, 10, r, 2) for r in range(1, 6)], ['3', '3', '3', '3', '12'])
        self.assertEqual(txt(d, 10, 5, 1), 'Total Score ( Max: 20 )')
        self.assertEqual([txt(d, 11, r, 4) for r in range(1, 7)], ['5', '5', '5', '5', '5', '25'])
        self.assertEqual([txt(d, 9, r, 1) for r in range(1, 6)], ['c0', 'c1', 'c2', 'c3', 'c4'])
        self.assertEqual(txt(d, 9, 7, 4), '12')
        t12 = [(txt(d, 12, r, 1), txt(d, 12, r, 3)) for r in range(len(d.tables[12].rows))]
        self.assertEqual(t12[2:7], [('e0', '4'), ('e1', '4'), ('e2', '4'), ('e3', '4'), ('Total (Max.20)', '16')])
        self.assertEqual(t12[-5:], [('p0', '2.5'), ('p1', '2.5'), ('p2', '2.5'), ('Total (Max.15)', '7.5'), (t12[-1][0], '23.5')])
        self.assertEqual([txt(d, 21, r, 0) for r in range(1, 7)], ['1', '2', '3', '4', '5', '6'])
        self.assertEqual(txt(d, 29, 11, 5), '18')
        self.assertEqual(txt(d, 23, 4, 3), '18')

    def test_place_line(self):
        def place(college):
            out = self.tmp / 'place.docx'
            generate_acr.generate({'college': college}, out)
            return [p.text.replace('	', '') for p in Document(out).paragraphs if p.text.startswith('Place: ') and 'reported on officer' in p.text][0]
        self.assertEqual(place({}), 'Place: Signature of the reported on officer')
        self.assertEqual(place({'name': 'Govt. College X'}), 'Place: Govt. College XSignature of the reported on officer')
        self.assertEqual(place({'pin': '171009', 'name': 'Govt. College X'}), 'Place: 171009, Govt. College XSignature of the reported on officer')

    def test_problems_block_generation(self):
        out = self.tmp / 'blocked.docx'
        with self.assertRaises(generate_acr.ApiProblemsError) as ctx:
            generate_acr.generate({'api': {'c1': {'classes': 60}}}, out)
        self.assertEqual(ctx.exception.problems[0]['where'], '26(i)(a)')
        self.assertFalse(out.exists())


if __name__ == '__main__':
    unittest.main()
