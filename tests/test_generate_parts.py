import json, shutil, sys, tempfile, unittest
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import generate_acr  # noqa: E402

FULL = json.loads((ROOT / 'tests' / 'fixtures' / 'full_record.acr.json').read_text(encoding='utf-8'))


class GenerateParts(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp())
        generate_acr.generate(FULL, cls.tmp / 'full.docx')
        cls.doc = Document(cls.tmp / 'full.docx')
        cls.paras = [p.text for p in cls.doc.paragraphs]
        # All lines in document order, including the text boxes of points 17, 18 and 19(b).
        def line(p):
            return ''.join('\n' if el.tag == qn('w:br') else '\t' if el.tag == qn('w:tab') else (el.text or '') if el.tag == qn('w:t') else ''
                           for el in p.iter())
        cls.text = '\n'.join(line(p) for p in cls.doc.element.body.iter(qn('w:p')) if p.find('.//' + qn('w:p')) is None)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def on_line(self, label, value):
        lines = [t for t in self.paras if label in t and value in t]
        self.assertTrue(lines, f'no line with "{label}" and "{value}"')

    def between(self, before, value, after):
        a = self.text.index(before)
        v = self.text.index(value, a)
        self.assertLess(v, self.text.index(after, a), value)

    def test_same_line_values(self):
        for label, value in [
            ('Name of the College through which ACR is submitted', 'Govt College Alpha'),
            ('Name of the College through which ACR is submitted', 'District Beta, 171001'),
            ('Submitted for the year/session', '2031-32'),
            ('Appraisal of work and conduct', 'ASHA DEVI'),
            ('Full Name (in Capital letter)', 'ASHA DEVI'),
            ('Father/Husband', 'RAM LAL'),
            ('Employee Code', 'EC-4321'),
            ('Subject for which Appointed', 'Chemistry'),
            ('Date of appointment', '01/07/2012'),
            ('Current Designation', 'Associate Professor'),
            ('Present Pay Band with Grade Pay', 'Level 13A; Basic Pay 131400'),
            ('Date of Promotion', 'No promotion'),
            ('Academic', 'M.Sc. Chemistry First Division'),
            ('Professional', 'NET'),
            ('Research Degree', 'Ph.D Chemistry'),
            ('In words', 'Fifth November Nineteen Hundred Eighty'),
            ('College/Colleges in which served', 'Govt College Alpha:\n\t01/04/2031 to 31/03/2032'),
            ('Roll no (with session)', 'Roll 12345, 2014'),
            ('b) Hindi subject : Cleared / exempted (mention details)', 'Cleared 2013'),
            ('Any other major assignment', 'Bursar'),
            ('Permanent Address (With Pin code)', 'House 12'),
            ('Land line telephone', '01972-222333'),
            ('Land line telephone', '+919800000001'),
            ('Email', 'asha@example.org'),
            ('Place: ', 'Govt College Alpha, 171001'),
            ('Designation, ', '15/05/2032'),
            ('Designation, ', 'Associate Professor'),
            ('Name of the Principal: ', 'Dr Principal Gamma'),
            ('(PBAS) FOR THE SESSION/YEAR', '2031-32'),
        ]:
            with self.subTest(label=label, value=value):
                self.on_line(label, value)
        self.assertTrue([t for t in self.paras if 'Permanent/Quasi-permanent' in t and t.strip().endswith('Permanent')])
        self.assertTrue([t for t in self.paras if 'Are you satisfied with your present position' in t and t.strip().endswith('No')])

    def test_answers_below_questions(self):
        self.between('Permanent Address (With Pin code)', 'Ward 3', 'Land line telephone')
        self.between('Permanent Address (With Pin code)', 'Town Delta, PIN 176001', 'Land line telephone')
        self.between('most important contribution', 'Contribution line 1\nContribution line 2', 'not assigned to you')
        self.between('not assigned to you', 'Unassigned work', 'Weekly time table')
        self.between('special effort made to improve class room', 'Special effort', 'How many assignments')
        self.between('Which new books', 'Book A - extract', 'vital problems')
        self.between('vital problems', 'Problem 1\nProblem 2', 'Details of Last year')
        self.between('fresh academic / professional qualifications', 'No fresh degree this year', 'Academic Staff College')
        self.between('Are you doing any Research work', 'Yes', 'Did you receive any honour')
        self.between('Did you receive any honour', 'Best teacher award', 'Are you satisfied')
        self.between('If not, do you want to change the profession', 'Want promotion', 'Any other significant point')
        self.between('Any other significant point', 'Other point', 'PART-II: SECTION-II')

    def test_tables(self):
        T = self.doc.tables
        cells = lambda t, r: [c.text.strip() for c in T[t].rows[r].cells]
        self.assertEqual(cells(0, 0), list('05111980'))
        self.assertEqual(cells(1, 1), ['1', 'B.Sc. I', 'GCA', '6', '150', '100%'])
        self.assertEqual(cells(1, 2), ['2', 'B.Sc. II', 'GCA', '6', '140', '95%'])
        self.assertEqual(cells(1, len(T[1].rows) - 1)[0], 'Total periods per week')
        self.assertEqual(cells(1, len(T[1].rows) - 1)[3], '24')
        self.assertEqual(cells(2, 1), ['1', 'B.Sc. I', '4', '2', ''])
        self.assertEqual(cells(3, 1), ['Chem Quiz', 'Inter-class quiz'])
        self.assertEqual(cells(4, 3), ['B.Sc. III', '1 year', '40', '38', '95', '88.5', '+6.5', '+6.5', '10', '20', '8', '2', ''])
        self.assertEqual(cells(5, 1), ['Refresher in Chemistry, UGC', 'HRDC Shimla', '21 days', 'RC-7'])
        self.assertEqual(cells(6, 1), ['Green synthesis', 'HPU', 'Minor', 'Ongoing'])
        self.assertEqual(cells(24, 1), ['1', 'Reviewer for journal X'])

    def test_no_token_left_and_enclosures(self):
        full = ' '.join(t.text or '' for t in self.doc.element.body.iter(qn('w:t')))
        self.assertNotIn('{{', full)
        i = self.paras.index('☑ 1. Certificate / sanction order')
        self.assertTrue(self.paras[i + 1].startswith('I certify that the information provided'))
        self.assertNotIn('Degree / qualification certificate', full)

    def strikes(self, doc, start, options):
        p = [p for p in doc.paragraphs if p.text.startswith(start)][0]
        return {r.text: bool(r.font.strike) for r in p.runs if r.text in options}

    def test_strike_through_choices(self):
        self.assertEqual(self.strikes(self.doc, 'Appraisal of work and conduct', ('Dr.', 'Shri', 'Smt', 'Kumari')),
                         {'Dr.': True, 'Shri': True, 'Smt': False, 'Kumari': True})
        self.assertEqual(self.strikes(self.doc, 'Father/Husband', ('Father', 'Husband')), {'Father': True, 'Husband': False})
        out = self.tmp / 'none.docx'
        generate_acr.generate({'profile': {}}, out)
        d = Document(out)
        self.assertFalse(any(self.strikes(d, 'Appraisal of work and conduct', ('Dr.', 'Shri', 'Smt', 'Kumari')).values()))
        self.assertFalse(any(self.strikes(d, 'Father/Husband', ('Father', 'Husband')).values()))

    def test_entry_problems_block(self):
        out = self.tmp / 'bad.docx'
        with self.assertRaises(generate_acr.ProblemsError) as ctx:
            generate_acr.generate({'profile': {'dob': '31/02/1990'}}, out)
        self.assertEqual(ctx.exception.problems[0]['code'], 'BAD_DOB')
        self.assertFalse(out.exists())


if __name__ == '__main__':
    unittest.main()
