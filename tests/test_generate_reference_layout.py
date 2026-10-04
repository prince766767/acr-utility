"""Output details matched to a teacher-made reference ACR: 26(i) rows, NIL / 00 in empty Category-III parts."""
import json, sys, tempfile, unittest
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import generate_acr  # noqa: E402

FULL = json.loads((ROOT / 'tests' / 'fixtures' / 'full_record.acr.json').read_text(encoding='utf-8'))
LECTURES = [{'course': 'BOTA-101', 'level': 'UG', 'mode': 'Traditional blackboard', 'allotted': '3 Theory + 1 Practical', 'conducted': '57', 'pct': '100%'},
            {'course': 'BOTA-301', 'level': 'UG', 'mode': 'Modern digital', 'allotted': '3 Theory', 'conducted': '61', 'pct': '100%'}]


def make(mutate):
    data = json.loads(json.dumps(FULL))
    mutate(data)
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / 'x.docx'
        generate_acr.generate(data, out)
        return Document(out)


def text(e):
    return ''.join(x.text or '' for x in e.iter(qn('w:t')))


def heading_before(doc, ti):
    prev = doc.tables[ti]._tbl.getprevious()
    while prev is not None and not (prev.tag == qn('w:p') and text(prev).strip()):
        prev = prev.getprevious()
    return prev


class Lectures26i(unittest.TestCase):
    def test_rows_are_printed_numbered(self):
        doc = make(lambda d: d.setdefault('api', {}).setdefault('c1', {}).update(lectures=LECTURES))
        t = doc.tables[7]
        self.assertTrue(t.rows[0].cells[1].text.startswith('Course'))
        self.assertEqual([c.text for c in t.rows[1].cells], ['1', 'BOTA-101', 'UG', 'Traditional blackboard', '3 Theory + 1 Practical', '57', '100%'])
        self.assertEqual(t.rows[2].cells[1].text, 'BOTA-301')
        self.assertEqual(len(t.rows), 6)  # header + the form's 5 lines

    def test_more_rows_than_lines_adds_rows(self):
        many = [dict(LECTURES[0], course=f'C{i}') for i in range(7)]
        doc = make(lambda d: d.setdefault('api', {}).setdefault('c1', {}).update(lectures=many))
        t = doc.tables[7]
        self.assertEqual(len(t.rows), 8)
        self.assertEqual(t.rows[7].cells[1].text, 'C6')


class NilForEmptyParts(unittest.TestCase):
    def test_empty_parts_say_nil_and_00(self):
        def empty_c3(d):
            d.setdefault('api', {})['c3'] = {}
        doc = make(empty_c3)
        for ti in range(13, 23):
            h = heading_before(doc, ti)
            self.assertTrue(text(h).rstrip().endswith('NIL'), (ti, text(h)))
            nil_run = h.findall(qn('w:r'))[-1]
            self.assertEqual(text(nil_run), 'NIL')
            self.assertGreaterEqual(len([r for r in h.findall(qn('w:r'))[-4:] if r.find(qn('w:tab')) is not None]), 1, ti)
        for ti in (13, 14, 15, 16, 17, 18, 20, 21, 22):
            self.assertEqual(doc.tables[ti].rows[1].cells[-1].text, '00', ti)
        self.assertEqual([doc.tables[19].rows[r].cells[4].text for r in (1, 2)], ['00', '00'])

    def test_filled_parts_have_no_nil(self):
        def one_paper(d):
            d.setdefault('api', {})['c3'] = {'journals': [{'title': 'Paper 1', 'journal': 'J', 'row': 'A1', 'score': '15'}],
                                             'guidance': {'phdAwarded': '1', 'phdAwardedScore': '10'}}
        doc = make(one_paper)
        self.assertFalse(text(heading_before(doc, 13)).rstrip().endswith('NIL'))
        self.assertEqual(doc.tables[13].rows[1].cells[-1].text, '15')
        self.assertFalse(text(heading_before(doc, 19)).rstrip().endswith('NIL'))
        self.assertTrue(text(heading_before(doc, 14)).rstrip().endswith('NIL'))


if __name__ == '__main__':
    unittest.main()


class Point12Lines(unittest.TestCase):
    def test_each_further_line_starts_with_a_tab(self):
        doc = make(lambda d: d.setdefault('part2', {}).update(p12='College A: 2021 to 2022\nCollege B: 2022 to date'))
        p = next(p for p in doc.element.body.iter(qn('w:p')) if text(p).startswith('College/Colleges in which served'))
        seq = [c.tag.split('}')[1] for r in p.findall(qn('w:r')) for c in r if c.tag in (qn('w:br'), qn('w:tab'))]
        # one TAB to the answer column, then BR + TAB before each of the 3 further lines (the template's own TAB follows)
        self.assertEqual(seq[:7], ['tab'] + ['br', 'tab'] * 3)
        self.assertIn('(1) College A:', text(p))


class InlineMarks(unittest.TestCase):
    def runs_of(self, doc, starts):
        p = next(p for p in doc.element.body.iter(qn('w:p')) if starts in text(p))
        out = []
        for r in p.findall(qn('w:r')):
            rpr, t = r.find(qn('w:rPr')), text(r)
            if t:
                out.append((t, rpr.find(qn('w:b')) is not None, rpr.find(qn('w:i')) is not None,
                            rpr.find(qn('w:vertAlign')) is not None))
        return out

    def test_bold_italic_superscript_become_runs(self):
        doc = make(lambda d: d['part2'].update(p17='Taught **Botany** to the 3^rd^ semester, *Biodiversity* too'))
        runs = self.runs_of(doc, 'Taught ')
        self.assertIn(('Botany', True, False, False), runs)
        self.assertIn(('rd', False, False, True), runs)
        self.assertIn(('Biodiversity', False, True, False), runs)
        self.assertNotIn('*', ''.join(t for t, *_ in runs))
        self.assertNotIn('^', ''.join(t for t, *_ in runs))

    def test_template_stars_untouched(self):
        doc = make(lambda d: None)
        self.assertIn('****************************', ''.join(text(p) for p in doc.element.body.iter(qn('w:p'))))
