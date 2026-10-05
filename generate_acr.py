from __future__ import annotations
import json, re, shutil, subprocess
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree
from docx import Document
from docx.shared import RGBColor, Inches, Pt
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from api_tally import tally, is_empty_entry, score_text, LEVEL_TEXT, last_year_problems, last_year_cells
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.run import Run
from acr_fields import token_values, part_tables, field_problems, inline_marks, TITLES, RELATIONS

HERE=Path(__file__).resolve().parent
TEMPLATE=HERE/'ACR_EMPLOYEE_MASTER.docx'
APPENDIX_DIR=HERE/'appendix_pages'
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
W='{%s}'%NS['w']
BLUE=RGBColor(0,0,204)


def set_cell(doc,ti,row,col,text):
    cell=doc.tables[ti].cell(row,col)
    if not cell.paragraphs: cell.add_paragraph()
    p=cell.paragraphs[0]
    if not p.runs: r=p.add_run()
    else: r=p.runs[0]
    r.text=str(text or '')
    r.font.color.rgb=BLUE
    give_cell_font(p,r)
    for rr in p.runs[1:]: rr.text=''


FORM_FONT='Times New Roman'
STYLE_FONTS=('Times New Roman','Arial','Calibri','Cambria','Georgia','Verdana')  # also in Google Docs, so the PDF matches
STYLE_SIZES=(10,11,12)
# Answers print black in the form's own fonts unless the teacher chose otherwise (Text style marks a choice "chosen").
DEFAULT_STYLE={'color':'000000','font':'','size':0,'bold':False,'italic':False,'chosen':False}
TEMPLATE_LOOK={'color':'0000CC','font':'','size':0,'bold':False,'italic':False}   # the blue answer runs as the template has them


def normalize_style(style):
    """The teacher's text style for filled-in answers. Not chosen (or saved before choosing existed): the default."""
    s=style if isinstance(style,dict) else {}
    if s.get('chosen') is not True:
        return dict(DEFAULT_STYLE)
    color=str(s.get('color') or '').lstrip('#').upper()
    try: size=int(s.get('size') or 0)
    except (TypeError,ValueError): size=0
    return {'color':color if re.fullmatch(r'[0-9A-F]{6}',color) else DEFAULT_STYLE['color'],
            'font':s.get('font') if s.get('font') in STYLE_FONTS else '',
            'size':size if size in STYLE_SIZES else 0,
            'bold':s.get('bold') is True,'italic':s.get('italic') is True,'chosen':True}


def apply_text_style(doc,style):
    """Every filled-in answer is a blue (0000CC) run: give each the teacher's colour, font, size, bold and italic."""
    st=normalize_style(style)
    if all(st[k]==v for k,v in TEMPLATE_LOOK.items()):
        return
    for r in list(doc.element.body.iter(qn('w:r'))):
        rpr=r.find(qn('w:rPr'))
        c=rpr.find(qn('w:color')) if rpr is not None else None
        if c is None or (c.get(qn('w:val')) or '').upper()!='0000CC':
            continue
        run=Run(r,None)
        if st['font']:
            fonts=r.get_or_add_rPr().get_or_add_rFonts()
            for k in list(fonts.attrib):
                if k.endswith('Theme'): del fonts.attrib[k]
            for k in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): fonts.set(qn(k),st['font'])
        if st['size']:
            run.font.size=Pt(st['size'])
            rpr=r.find(qn('w:rPr'))
            szcs=rpr.find(qn('w:szCs'))
            if szcs is None:
                from docx.oxml import OxmlElement
                szcs=OxmlElement('w:szCs'); rpr.find(qn('w:sz')).addnext(szcs)
            szcs.set(qn('w:val'),str(st['size']*2))
        if st['bold']:
            run.font.bold=True; run.font.cs_bold=True
        if st['italic']:
            run.font.italic=True; run.font.cs_italic=True
        c.set(qn('w:val'),st['color'])
    # Paragraph marks of answers (they colour automatic numbers): colour only, their size would change line heights.
    for c in doc.element.body.iter(qn('w:color')):
        if c.getparent().getparent().tag==qn('w:pPr') and (c.get(qn('w:val')) or '').upper()=='0000CC':
            c.set(qn('w:val'),st['color'])


def give_cell_font(p,r):
    """A run with no font of its own falls back to the template default (Calibri 11). Give it the font and size set
    on the cell's paragraph mark, or the form's Times New Roman when the mark has none."""
    if r._r.rPr is not None and r._r.rPr.find(qn('w:rFonts')) is not None:
        return
    ppr=p._p.pPr
    mark=ppr.find(qn('w:rPr')) if ppr is not None else None
    fonts=mark.find(qn('w:rFonts')) if mark is not None else None
    rpr=r._r.get_or_add_rPr()
    rfonts=rpr.get_or_add_rFonts()  # python-docx puts it in its schema position
    if fonts is not None and fonts.get(qn('w:ascii')):
        for k,v in fonts.attrib.items(): rfonts.set(k,v)
    else:
        for k in ('w:ascii','w:hAnsi','w:cs'): rfonts.set(qn(k),FORM_FONT)
    sz=mark.find(qn('w:sz')) if mark is not None else None
    if sz is not None and rpr.find(qn('w:sz')) is None:
        r.font.size=Pt(int(sz.get(qn('w:val')))/2)
        szcs=mark.find(qn('w:szCs'))
        if szcs is not None and rpr.find(qn('w:szCs')) is None:
            rpr.find(qn('w:sz')).addnext(deepcopy(szcs))  # szCs follows sz in the schema


def clone_row(table,row_idx):
    tr=table.rows[row_idx]._tr
    new_tr=deepcopy(tr)
    tr.addnext(new_tr)
    return table.rows[row_idx+1]


def fill_rows(doc,ti,rows,start=1):
    t=doc.tables[ti]
    existing=max(0,len(t.rows)-start)
    if len(rows)>existing:
        for _ in range(len(rows)-existing): clone_row(t,start)
    for i,row in enumerate(rows):
        vals=row if isinstance(row,list) else [row.get(str(c),'') for c in range(len(t.columns))]
        for c in range(len(t.columns)): set_cell(doc,ti,start+i,c, vals[c] if c<len(vals) else '')


def merge_down(doc,ti,col,first,last):
    """One merged cell (Word vMerge) in column col over rows first..last; its text is the first row's. As mergeDown in docx_engine.js."""
    t=doc.tables[ti]
    tcs=[t.cell(r,col)._tc for r in range(first,last+1)]   # all found before any is merged
    for k,tc in enumerate(tcs):
        tcpr=tc.get_or_add_tcPr()
        v=OxmlElement('w:vMerge')
        v.set(qn('w:val'),'restart' if k==0 else 'continue')
        anchor=next((e for e in (tcpr.find(qn('w:gridSpan')),tcpr.find(qn('w:tcW'))) if e is not None),None)
        if anchor is not None: anchor.addnext(v)
        else: tcpr.insert(0,v)


class ProblemsError(Exception):
    """The teacher's entries have problems; the ACR must not be generated."""
    def __init__(self, problems):
        self.problems = problems
        super().__init__('\n'.join(p['message'] for p in problems))


ApiProblemsError = ProblemsError  # earlier name, kept for callers


# Point 44 "API Score reported in self appraisal" cells: code -> (table, row, column, start of Max. Score text).
# The Max. Score column is the column just before the API column; it is checked before writing.
P44_CELLS = {
    'A1': (27, 3, 4, '15/'), 'A2': (27, 4, 4, '10/'), 'B1a': (27, 5, 4, '10/'), 'B1b': (27, 6, 4, '5/'),
    'B2': (28, 0, 4, '10/'), 'B3a': (28, 1, 4, '50/'), 'B3b': (28, 2, 4, '25/'), 'B3c': (28, 3, 4, '15/'),
    'C1a': (28, 4, 4, '20/'), 'C1b': (28, 5, 4, '15/'), 'C1c': (28, 6, 4, '10/'),
    'C2': (29, 0, 5, '10/'), 'C3': (29, 1, 5, '20/'), 'C4': (29, 2, 5, '30/'),
    'D1': (29, 3, 5, '3/'), 'D2a': (29, 4, 5, '10/'), 'D2b': (29, 5, 5, '7/'),
    'E1a': (29, 6, 5, '20/'), 'E1b': (29, 7, 5, '10/'),
    'E2a': (29, 8, 5, '10/'), 'E2b': (29, 9, 5, '7.5/'), 'E2c': (29, 10, 5, '5/'), 'E2d': (29, 11, 5, '3/'),
    'E3a': (30, 0, 4, '10/'), 'E3b': (30, 1, 4, '5/'),
}


def _norm(s):
    return re.sub(r'\s+', '', s or '')


def _check(ok, ti, what):
    if not ok:
        raise RuntimeError(f'Template table {ti}: expected {what}. The template layout has changed; nothing was written.')


def find_row(table, col, prefix, start=0):
    for r in range(start, len(table.rows)):
        if _norm(table.rows[r].cells[col].text).startswith(_norm(prefix)):
            return r
    raise RuntimeError(f'Template: no row starting with "{prefix}" in column {col}.')


def fill_between(doc, ti, first, end, rows):
    """Write rows into table ti from row `first`, before row `end`; clone blank entry rows when more are needed."""
    t = doc.tables[ti]
    free = end - first
    for _ in range(len(rows) - free):
        clone_row(t, end - 1)
    for i, vals in enumerate(rows):
        for c, v in enumerate(vals):
            set_cell(doc, ti, first + i, c, v)


def entries(api, group, name):
    g = api.get(group) if isinstance(api.get(group), dict) else {}
    lst = g.get(name) if isinstance(g.get(name), list) else []
    return [e for e in lst if not is_empty_entry(e)]


def field(e, key):
    v = e.get(key)
    return '' if v is None else str(v)


# Category-III parts: table, list name, words in the heading just above the table (checked before writing).
C3_PARTS = ((13, 'journals', 'Publishedpapers'), (14, 'chapters', 'Articles/Chapters'), (15, 'proceedings', 'FullPapers'),
            (16, 'books', 'BooksPublished'), (17, 'ongoing', 'OngoingProjects'), (18, 'completed', 'CompletedProjects'),
            (20, 'training', 'TrainingCourses'), (21, 'papers', 'Paperspresented'), (22, 'lectures', 'InvitedLectures'))
GUIDANCE_KEYS = ('mphilEnrolled', 'mphilSubmitted', 'mphilAwarded', 'mphilScore', 'phdEnrolled', 'phdSubmitted', 'phdAwarded',
                 'phdAwardedScore', 'phdSubmittedScore')


def heading_above(doc, ti, words):
    prev = doc.tables[ti]._tbl.getprevious()
    while prev is not None and not (prev.tag == qn('w:p') and ''.join(x.text or '' for x in prev.iter(qn('w:t'))).strip()):
        prev = prev.getprevious()
    _check(prev is not None and words in _norm(''.join(x.text or '' for x in prev.iter(qn('w:t')))), ti, f'the heading "{words}" above it')
    return prev


def mark_nil(doc, ti, words):
    """An empty Category-III part: "NIL" after its heading (two tabs on), as teachers write it."""
    p = heading_above(doc, ti, words)
    last = [r for r in p.findall(qn('w:r')) if r.find(qn('w:t')) is not None][-1]
    rpr = last.find(qn('w:rPr'))
    tabs = etree.SubElement(p, W + 'r')
    if rpr is not None: tabs.append(deepcopy(rpr))
    etree.SubElement(tabs, W + 'tab'); etree.SubElement(tabs, W + 'tab')
    nil = etree.SubElement(p, W + 'r')
    if rpr is not None: nil.append(deepcopy(rpr))
    Run(nil, None).font.color.rgb = BLUE
    etree.SubElement(nil, W + 't').text = 'NIL'


def fill_api_tables(doc, api, v):
    T = doc.tables
    # 26(i) lectures, seminars, tutorials, practicals
    _check(_norm(T[7].rows[0].cells[1].text).startswith('Course') and len(T[7].columns) == 7, 7, 'the 26(i) Course/Paper table')
    fill_rows(doc, 7, [[i, field(e, 'course'), field(e, 'level'), field(e, 'mode'), field(e, 'allotted'), field(e, 'conducted'), field(e, 'pct')]
                       for i, e in enumerate(entries(api, 'c1', 'lectures'), 1)], 1)
    # 26(i)(a), (b)
    _check(_norm(T[8].rows[1].cells[0].text) == '(a)' and _norm(T[8].rows[2].cells[0].text) == '(b)', 8, '(a)/(b) rows')
    set_cell(doc, 8, 1, 2, v['p42']['i_a'])
    set_cell(doc, 8, 2, 2, v['p42']['i_b'])
    # 26(ii)
    fill_between(doc, 9, 1, find_row(T[9], 0, 'API score based'),
                 [[i, field(e, 'course'), field(e, 'consulted'), field(e, 'prescribed'), field(e, 'additional')]
                  for i, e in enumerate(entries(api, 'c1', 'resources'), 1)])
    score_row = find_row(T[9], 0, 'APIscorebased') + 1
    _check(_norm(T[9].rows[score_row - 1].cells[4].text) == 'APIScore', 9, '"API Score" above the 26(ii) score cell')
    set_cell(doc, 9, score_row, 4, v['p42']['ii'])
    # 26(iii)
    fill_between(doc, 10, 1, find_row(T[10], 1, 'Total Score'),
                 [[i, field(e, 'description'), score_text(e.get('score'))] for i, e in enumerate(entries(api, 'c1', 'innovative'), 1)])
    set_cell(doc, 10, find_row(T[10], 1, 'Total Score'), 2, v['p42']['iii'])
    # 26(iv)
    fill_between(doc, 11, 1, find_row(T[11], 1, 'Total Score'),
                 [[i, field(e, 'type'), field(e, 'assigned'), field(e, 'extent'), score_text(e.get('score'))]
                  for i, e in enumerate(entries(api, 'c1', 'exam'), 1)])
    set_cell(doc, 11, find_row(T[11], 1, 'Total Score'), 4, v['p42']['iv'])
    # 27 (i), (ii), (iii)
    for heading, total_label, name, col2, total in (('(i)', 'Total (Max.20)', 'extension', 'hours', v['p43']['i']),
                                                    ('(ii)', 'Total (Max.15)', 'management', 'responsibility', v['p43']['ii']),
                                                    ('(iii)', 'Total (Max.15)', 'professional', 'details', v['p43']['iii'])):
        h = find_row(T[12], 1, heading)
        end = find_row(T[12], 1, total_label, h + 1)
        fill_between(doc, 12, h + 1, end,
                     [[i, field(e, 'activity'), field(e, col2), score_text(e.get('score'))] for i, e in enumerate(entries(api, 'c2', name), 1)])
        set_cell(doc, 12, find_row(T[12], 1, total_label, h + 1), 3, total)
    set_cell(doc, 12, find_row(T[12], 1, 'Total Score'), 3, v['p43']['total'])
    # 28 A ... E(iii) detail tables
    sc = lambda e: score_text(e.get('score'))
    fill_rows(doc, 13, [[i, field(e, 'title'), field(e, 'journal'), field(e, 'issn'), field(e, 'peer'), field(e, 'coauthors'), field(e, 'mainAuthor'), sc(e)]
                        for i, e in enumerate(entries(api, 'c3', 'journals'), 1)], 1)
    fill_rows(doc, 14, [[i, field(e, 'title'), field(e, 'book'), field(e, 'issn'), field(e, 'peer'), field(e, 'coauthors'), field(e, 'mainAuthor'), sc(e)]
                        for i, e in enumerate(entries(api, 'c3', 'chapters'), 1)], 1)
    fill_rows(doc, 15, [[i, field(e, 'title'), field(e, 'conference'), field(e, 'issn'), field(e, 'coauthors'), field(e, 'mainAuthor'), sc(e)]
                        for i, e in enumerate(entries(api, 'c3', 'proceedings'), 1)], 1)
    fill_rows(doc, 16, [[i, field(e, 'title'), field(e, 'type'), field(e, 'publisher'), field(e, 'peer'), field(e, 'coauthors'), field(e, 'mainAuthor'), sc(e)]
                        for i, e in enumerate(entries(api, 'c3', 'books'), 1)], 1)
    fill_rows(doc, 17, [[i, field(e, 'title'), field(e, 'agency'), field(e, 'period'), field(e, 'amount'), sc(e)]
                        for i, e in enumerate(entries(api, 'c3', 'ongoing'), 1)], 1)
    fill_rows(doc, 18, [[i, field(e, 'title'), field(e, 'agency'), field(e, 'period'), field(e, 'amount'), field(e, 'outcome'), sc(e)]
                        for i, e in enumerate(entries(api, 'c3', 'completed'), 1)], 1)
    g = api.get('c3', {}).get('guidance', {}) if isinstance(api.get('c3'), dict) else {}
    g = g if isinstance(g, dict) else {}
    _check(T[19].rows[1].cells[0].text.startswith('M.Phil') and T[19].rows[2].cells[0].text.startswith('Ph.D'), 19, 'M.Phil / Ph.D rows')
    for c, key in ((1, 'mphilEnrolled'), (2, 'mphilSubmitted'), (3, 'mphilAwarded')):
        set_cell(doc, 19, 1, c, field(g, key))
    set_cell(doc, 19, 1, 4, v['p44']['D1'])
    for c, key in ((1, 'phdEnrolled'), (2, 'phdSubmitted'), (3, 'phdAwarded')):
        set_cell(doc, 19, 2, c, field(g, key))
    set_cell(doc, 19, 2, 4, v['p28']['phd'])
    fill_rows(doc, 20, [[i, field(e, 'programme'), field(e, 'duration'), field(e, 'organisedBy'), sc(e)]
                        for i, e in enumerate(entries(api, 'c3', 'training'), 1)], 1)
    fill_rows(doc, 21, [[i, field(e, 'title'), field(e, 'conference'), field(e, 'organisedBy'), LEVEL_TEXT.get(e.get('row'), ''), sc(e)]
                        for i, e in enumerate(entries(api, 'c3', 'papers'), 1)], 1)
    fill_rows(doc, 22, [[i, field(e, 'title'), field(e, 'conference'), field(e, 'organisedBy'), LEVEL_TEXT.get(e.get('row'), ''), sc(e)]
                        for i, e in enumerate(entries(api, 'c3', 'lectures'), 1)], 1)
    # Empty Category-III parts: NIL by the heading, 00 in the score box.
    for ti, name, words in C3_PARTS:
        if not entries(api, 'c3', name):
            mark_nil(doc, ti, words)
            set_cell(doc, ti, 1, len(T[ti].rows[1].cells) - 1, '00')
    if not any(field(g, k).strip() for k in GUIDANCE_KEYS):
        mark_nil(doc, 19, 'ResearchGuidance')
        set_cell(doc, 19, 1, 4, '00')
        set_cell(doc, 19, 2, 4, '00')
    # 29 and 45: column 3 = last academic year (I+II computed), column 4 = this year's totals; 45 col 5 is the Principal's.
    ly = last_year_cells(api.get('lastAcademicYear'))
    _check(_norm(T[31].rows[0].cells[2].text).startswith('LastAcademic'), 31, 'the "Last Academic Year" header of point 45')
    for r, label, key, value in ((1, 'Teaching', 'cat1', v['p29']['I']), (2, 'Co-curricular', 'cat2', v['p29']['II']),
                                 (3, 'Total', 'total12', v['p29']['I_II']), (4, 'Research', 'cat3', v['p29']['III'])):
        for ti in (23, 31):
            _check(_norm(T[ti].rows[r].cells[1].text).startswith(label), ti, f'row {r} starting "{label}"')
            set_cell(doc, ti, r, 2, ly[key])
            set_cell(doc, ti, r, 3, value)
    # 42: column 4
    for r, label, key in ((3, '(i)a', 'i_a'), (4, '(i)b', 'i_b'), (5, '(ii)', 'ii'), (6, '(iii)', 'iii'), (7, '(iv)', 'iv')):
        _check(_norm(T[25].rows[r].cells[0].text) == label, 25, f'row {r} "{label}"')
        set_cell(doc, 25, r, 3, v['p42'][key])
    _check(_norm(T[25].rows[8].cells[1].text).startswith('TotalScore'), 25, 'Total Score row')
    set_cell(doc, 25, 8, 3, v['p42']['total'])
    # 43: column 4
    for r, label, key in ((3, '(i)', 'i'), (4, '(ii)', 'ii'), (5, '(iii)', 'iii')):
        _check(_norm(T[26].rows[r].cells[0].text) == label, 26, f'row {r} "{label}"')
        set_cell(doc, 26, r, 3, v['p43'][key])
    _check(_norm(T[26].rows[6].cells[1].text).startswith('TotalScore'), 26, 'Total Score row')
    set_cell(doc, 26, 6, 3, v['p43']['total'])
    # 44: column 5, one cell per sub-row, then Total
    for code, (ti, r, c, max_prefix) in P44_CELLS.items():
        _check(_norm(T[ti].rows[r].cells[c - 1].text).startswith(max_prefix), ti, f'row {r} Max. Score "{max_prefix}" for {code}')
        _check(T[ti].rows[r].cells[c].text.strip() == '', ti, f'empty API cell for {code}')
        set_cell(doc, ti, r, c, v['p44'][code])
    _check(_norm(T[30].rows[2].cells[1].text) == 'Total', 30, 'Total row')
    set_cell(doc, 30, 2, 4, v['p44']['total'])


def replace_tokens(docx_path, out_path, values):
    with ZipFile(docx_path,'r') as zin: files={n:zin.read(n) for n in zin.namelist()}
    root=etree.fromstring(files['word/document.xml'])
    token_map={'{{'+k+'}}':str(v or '') for k,v in values.items()}
    # First pass: replace tokens within individual text nodes, preserving all
    # surrounding tabs, spacing and run structure in the official template.
    replaced=set()
    for tnode in root.xpath('.//w:t',namespaces=NS):
        txt=tnode.text or ''
        newtxt=txt
        for tok,val in token_map.items():
            if tok in newtxt:
                newtxt=newtxt.replace(tok,val); replaced.add(tok)
        if newtxt!=txt: tnode.text=newtxt
    # Multi-line answers: turn each newline into a Word line break inside the same run.
    for tnode in root.xpath('.//w:t',namespaces=NS):
        if tnode.text and '\n' in tnode.text:
            parts=tnode.text.split('\n')
            tnode.text=parts[0]; tnode.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
            prev=tnode
            for part in parts[1:]:
                br=etree.Element(W+'br'); prev.addnext(br)
                t=etree.Element(W+'t'); t.set('{http://www.w3.org/XML/1998/namespace}space','preserve'); t.text=part
                br.addnext(t); prev=t
    # Second pass: handle any token that happened to be split across XML runs.
    for p in root.xpath('.//w:p',namespaces=NS):
        full=''.join(p.xpath('.//w:t/text()',namespaces=NS))
        if not any(tok in full for tok in token_map):
            continue
        changed=full
        hit=False
        for tok,val in token_map.items():
            if tok in changed:
                changed=changed.replace(tok,val); hit=True
        if not hit: continue
        texts=p.xpath('.//w:t',namespaces=NS)
        if not texts:
            r=etree.SubElement(p,W+'r'); t=etree.SubElement(r,W+'t'); t.text=changed; continue
        texts[0].text=changed
        for tt in texts[1:]: tt.text=''
    files['word/document.xml']=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone='yes')
    with ZipFile(out_path,'w',ZIP_DEFLATED) as zout:
        for n,b in files.items(): zout.writestr(n,b)


def fill_part_tables(doc, t):
    """Point 10 digit boxes and the tables of points 19(a), 19(c), 19(d), 20, 21(ii), 22 and 30."""
    T = doc.tables
    _check(len(T[0].rows) == 1 and len(T[0].rows[0].cells) == 8, 0, 'the 8 date-of-birth boxes')
    for c, ch in enumerate(t['dob_digits'] or ' ' * 8):
        set_cell(doc, 0, 0, c, ch.strip())
    _check(_norm(T[1].rows[0].cells[0].text).startswith('Sr.'), 1, 'the 19(a) header')
    fill_between(doc, 1, 1, find_row(T[1], 0, 'Total periods per week'), t['teaching'])
    set_cell(doc, 1, find_row(T[1], 0, 'Total periods per week'), 3, t['total_periods'])
    _check(_norm(T[2].rows[0].cells[0].text).startswith('Sr.'), 2, 'the 19(c) header')
    fill_rows(doc, 2, t['assignments'], 1)
    if t['assignments_merged']:
        merge_down(doc, 2, 4, 1, len(t['assignments']))
    _check(_norm(T[3].rows[0].cells[0].text).startswith('Titleoftheactivity'), 3, 'the 19(d) header')
    fill_rows(doc, 3, t['activities'], 1)
    _check(_norm(T[4].rows[2].cells[0].text) == '1', 4, 'the column-number row of point 20')
    fill_rows(doc, 4, t['results'], 3)
    _check(_norm(T[5].rows[0].cells[0].text).startswith('NameoftheSummer'), 5, 'the 21(ii) header')
    fill_rows(doc, 5, t['orientation'], 1)
    _check(_norm(T[6].rows[0].cells[0].text).startswith('Topictitle'), 6, 'the 22 header')
    fill_rows(doc, 6, t['research'], 1)
    _check(_norm(T[24].rows[0].cells[0].text) == 'S.No.', 24, 'the point 30 header')
    fill_rows(doc, 24, t['other_info'], 1)


def strike_unchosen(doc, starts_with, options, chosen):
    """Strike through the options not chosen (cover-page title, point 2 Father/Husband). Nothing chosen: none struck."""
    for p in doc.element.body.iter(qn('w:p')):
        if ''.join(x.text or '' for x in p.iter(qn('w:t'))).startswith(starts_with):
            runs = {''.join(x.text or '' for x in r.iter(qn('w:t'))): r for r in p.findall(qn('w:r'))}
            if not all(o in runs for o in options):
                raise RuntimeError(f'Template: "{starts_with}" line has no separate runs for {options}.')
            for o in options:
                Run(runs[o], None).font.strike = chosen in options and o != chosen
            return
    raise RuntimeError(f'Template: no line starting with "{starts_with}".')


def write_run_text(r,text):
    """Run content as text, tabs and line breaks (as setRunText in docx_engine.js)."""
    for ch in [c for c in r if c.tag!=qn('w:rPr')]: r.remove(ch)
    for part in re.split(r'(\t|\n)',text):
        if part=='\t': etree.SubElement(r,W+'tab')
        elif part=='\n': etree.SubElement(r,W+'br')
        elif part:
            t=etree.SubElement(r,W+'t'); t.text=part
            if part.strip()!=part: t.set('{http://www.w3.org/XML/1998/namespace}space','preserve')


def apply_inline_marks(doc):
    """Answers (blue runs) marked **bold**, *italic* or ^superscript^: split into runs with that formatting."""
    kinds={qn('w:t'):None,qn('w:tab'):'\t',qn('w:br'):'\n'}
    for r in list(doc.element.body.iter(qn('w:r'))):
        rpr=r.find(qn('w:rPr'))
        c=rpr.find(qn('w:color')) if rpr is not None else None
        if c is None or (c.get(qn('w:val')) or '').upper()!='0000CC': continue
        content=[ch for ch in r if ch.tag!=qn('w:rPr')]
        if any(ch.tag not in kinds for ch in content): continue
        text=''.join((ch.text or '') if ch.tag==qn('w:t') else kinds[ch.tag] for ch in content)
        if '*' not in text and '^' not in text: continue
        segs=inline_marks(text)
        if all(not fmt for _,fmt in segs): continue
        for s,fmt in segs:
            nr=deepcopy(r)
            f=Run(nr,None).font
            if 'b' in fmt: f.bold=True; f.cs_bold=True
            if 'i' in fmt: f.italic=True; f.cs_italic=True
            if 's' in fmt: f.superscript=True
            write_run_text(nr,s)
            r.addprevious(nr)
        r.getparent().remove(r)


def indent_point12(doc):
    """Point 12: every further line of the answer starts with a TAB, so it lands on the answer-column tab stop."""
    for p in doc.element.body.iter(qn('w:p')):
        if ''.join(x.text or '' for x in p.iter(qn('w:t'))).startswith('College/Colleges in which served'):
            for br in list(p.iter(qn('w:br'))):
                if br.get(qn('w:type')) in (None, 'textWrapping'):
                    br.addnext(etree.Element(W+'tab'))
            return
    raise RuntimeError('Template: no line starting with "College/Colleges in which served".')


def generate(data,out_docx):
    api=data.get('api') if isinstance(data.get('api'),dict) else {}
    result=tally(api)
    problems=field_problems(data)+result['problems']+last_year_problems(api.get('lastAcademicYear'))
    if problems:
        raise ProblemsError(problems)
    v=result['values']
    profile=data.get('profile') if isinstance(data.get('profile'),dict) else {}
    doc=Document(TEMPLATE)
    fill_part_tables(doc,part_tables(data))
    strike_unchosen(doc,'Appraisal of work and conduct',TITLES,profile.get('title'))
    strike_unchosen(doc,'Father/Husband',RELATIONS,profile.get('relation'))
    fill_api_tables(doc,api,v)
    # Replace tokens in XML after table edits.
    tmp=out_docx.with_suffix('.intermediate.docx')
    doc.save(tmp)
    replace_tokens(tmp,out_docx,token_values(data))
    tmp.unlink(missing_ok=True)
    # Enclosures: the checked items, numbered, just before the teacher's certificate.
    doc=Document(out_docx)
    indent_point12(doc)
    selected=[x.get('label') for x in data.get('enclosures',[]) if isinstance(x,dict) and x.get('checked',True)]
    if selected:
        insert_at=None
        for p in doc.paragraphs:
            if p.text.strip().startswith('I certify that the information provided'):
                insert_at=p
                break
        if insert_at is not None:
            for idx,label in enumerate(selected,1):
                np=insert_at.insert_paragraph_before(f'☑ {idx}. {label}')
                for r in np.runs: r.font.color.rgb=BLUE
    apply_inline_marks(doc)
    apply_text_style(doc,data.get('style'))
    doc.save(out_docx)
    return v

if __name__=='__main__':
    import argparse
    ap=argparse.ArgumentParser(); ap.add_argument('json'); ap.add_argument('-o','--output',default='ACR_generated_v03.docx'); args=ap.parse_args()
    d=json.loads(Path(args.json).read_text(encoding='utf-8'))
    try:
        values=generate(d,Path(args.output))
    except ProblemsError as err:
        import sys
        print('ACR not generated. Fix these entries first:',file=sys.stderr)
        for prob in err.problems: print(' - '+prob['message'],file=sys.stderr)
        sys.exit(1)
    print(json.dumps(values,indent=2))
